import streamlit as st
import pandas as pd
from streamlit_gsheets import GSheetsConnection
from datetime import datetime, date
import uuid
import io
import base64
import time
import json
import os
from PIL import Image, ImageDraw
from shared_employees import employee_names, employee_selector
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph
from reportlab.lib import colors
from reportlab.lib.units import mm
from xml.sax.saxutils import escape

try:
    from streamlit_drawable_canvas import st_canvas
except ImportError:
    st_canvas = None

SIGNATURE_COLUMNS = ['customer_signature', 'signed_by', 'signed_at']
FORM_COLUMNS = ['address', 'manufacture_date', 'arrival_time', 'departure_time',
                'is_completed', 'next_service_date', 'next_service_hours', 'service_fee',
                'tax_amount', 'service_supervisor', 'sales_supervisor'] + [
    f'part_{n}_{field}' for n in range(1, 11) for field in ('name', 'price', 'qty')]
SHEET = 'ServiceTickets'
COLUMNS = ['ticket_id', 'created_at', 'customer', 'contact', 'phone', 'machine_model',
           'serial_number', 'warranty', 'issue', 'priority', 'status', 'engineer',
           'scheduled_date', 'diagnosis', 'repair_action', 'parts', 'parts_cost',
           'labor_hours', 'labor_rate', 'total_cost', 'resolution', 'closed_at',
           'updated_at'] + SIGNATURE_COLUMNS + FORM_COLUMNS
STATUSES = ['新報修', '待確認', '待派工', '待零件', '維修中', '待驗收', '已結案']
PRIORITIES = ['一般', '急件', '緊急']
WARRANTIES = ['待確認', '保固內', '保固外']

st.set_page_config(page_title='JENN-WEI 售服維修管理系統', page_icon='🛠️', layout='wide')
st.title('🛠️ JENN-WEI 售服維修管理系統')
st.caption('Service Management | 報修、派工、維修、驗收與售服績效')
nav_aa, nav_bb, _ = st.columns([1.5, 1.5, 4])
with nav_aa:
    st.link_button('📋 返回 RFQ 詢價管理', 'https://aazzyyii.streamlit.app/', use_container_width=True)
with nav_bb:
    st.link_button('📊 返回 Trello 任務管理', 'https://aazzyybb.streamlit.app/', use_container_width=True)


def get_spreadsheet_url():
    try:
        direct = st.secrets.get('SERVICE_SPREADSHEET_URL', '')
        settings = st.secrets.get('connections', {}).get('gsheets', {})
        return str(direct or settings.get('spreadsheet', '') or settings.get('spreadsheet_url', '')).strip()
    except Exception:
        return ''


SPREADSHEET_URL = get_spreadsheet_url()
if not SPREADSHEET_URL:
    st.error('尚未設定售服系統的 Google Spreadsheet 網址。')
    st.info('請在 Streamlit Cloud → Settings → Secrets 設定 SERVICE_SPREADSHEET_URL，並保留 [connections.gsheets] 服務帳號憑證。')
    st.stop()
conn = st.connection('gsheets', type=GSheetsConnection)


@st.cache_data(ttl=600, show_spinner=False)
def _cached_ticket_read(spreadsheet_url, _connection):
    for attempt in range(3):
        try:
            return _connection.read(spreadsheet=spreadsheet_url, worksheet=SHEET, ttl=600)
        except Exception as exc:
            quota = any(k in str(exc).lower() for k in ('429', 'resource_exhausted', 'quota exceeded', 'rate_limit_exceeded'))
            if not quota or attempt == 2:
                raise
            time.sleep(5 * 2 ** attempt)


def read_tickets():
    try:
        data = _cached_ticket_read(SPREADSHEET_URL, conn)
        data = pd.DataFrame() if data is None else data.copy()
        for col in COLUMNS:
            if col not in data.columns:
                data[col] = ''
        return data[COLUMNS].astype(object)
    except Exception as exc:
        if any(k in str(exc).lower() for k in ('429', 'resource_exhausted', 'quota exceeded', 'rate_limit_exceeded')):
            st.error('Google Sheets API 讀取配額已達上限（429）。請稍候約 1～2 分鐘再重新整理。')
        else:
            st.error(f'無法讀取 {SHEET} 工作表。請檢查試算表網址、工作表及服務帳號權限。')
        st.exception(exc)
        st.stop()


def save_tickets(data):
    try:
        conn.update(spreadsheet=SPREADSHEET_URL, worksheet=SHEET,
                    data=data[COLUMNS].copy().astype(object).fillna(''))
        _cached_ticket_read.clear()
    except Exception as exc:
        st.error('Google Sheets 寫入失敗')
        st.exception(exc)
        return
    st.success('資料已儲存')
    st.rerun()


def number(value):
    try:
        result = pd.to_numeric(value, errors='coerce')
        return 0.0 if pd.isna(result) else float(result)
    except (TypeError, ValueError):
        return 0.0


def safe_str(value):
    return '' if value is None or pd.isna(value) else str(value)


def now_text():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def signature_from_drawing(drawing, width=340, height=170):
    if isinstance(drawing, str):
        drawing = json.loads(drawing)
    if not isinstance(drawing, dict):
        raise ValueError('簽名資料格式異常，請重新簽名。')
    image = Image.new('RGB', (width * 4, height * 4), 'white')
    pen = ImageDraw.Draw(image)
    count = 0
    for obj in drawing.get('objects', []):
        if not isinstance(obj, dict):
            continue
        commands = obj.get('path', [])
        if isinstance(commands, str):
            commands = json.loads(commands)
        offset = obj.get('pathOffset') or {}
        w, h = float(obj.get('width') or 0), float(obj.get('height') or 0)
        ox, oy = float(offset.get('x', w / 2)), float(offset.get('y', h / 2))
        left, top = float(obj.get('left') or 0), float(obj.get('top') or 0)
        sx, sy = float(obj.get('scaleX') or 1), float(obj.get('scaleY') or 1)
        stroke = max(2, round(float(obj.get('strokeWidth') or 3)))

        def pos(x, y):
            return left + (float(x) - ox + w / 2) * sx, top + (float(y) - oy + h / 2) * sy

        previous = None
        for cmd in commands:
            if not isinstance(cmd, (list, tuple)) or not cmd:
                continue
            kind, args = str(cmd[0]).upper(), cmd[1:]
            if kind == 'M' and len(args) >= 2:
                previous = pos(*args[:2])
                x, y = previous
                pen.ellipse((x - 2, y - 2, x + 2, y + 2), fill='black')
                count += 1
            elif kind == 'L' and len(args) >= 2 and previous is not None:
                point = pos(*args[:2])
                pen.line([previous, point], fill='black', width=stroke)
                previous = point
                count += 1
            elif kind in ('Q', 'C') and previous is not None:
                required = 4 if kind == 'Q' else 6
                if len(args) < required:
                    continue
                pts = [pos(args[i], args[i + 1]) for i in range(0, required, 2)]
                segment = []
                for i in range(31):
                    t, u = i / 30, 1 - i / 30
                    if kind == 'Q':
                        point = tuple(u*u*previous[a] + 2*u*t*pts[0][a] + t*t*pts[1][a] for a in (0, 1))
                    else:
                        point = tuple(u**3*previous[a] + 3*u*u*t*pts[0][a] + 3*u*t*t*pts[1][a] + t**3*pts[2][a] for a in (0, 1))
                    segment.append(point)
                pen.line(segment, fill='black', width=stroke)
                previous = pts[-1]
                count += 1
    bounds = Image.eval(image.convert('L'), lambda v: 255 - v).getbbox()
    if not count or bounds is None:
        raise ValueError('尚未偵測到簽名筆跡，請重新簽名或上傳簽名圖片。')
    cropped = image.crop((max(0, bounds[0]-5), max(0, bounds[1]-5),
                          min(image.width, bounds[2]+5), min(image.height, bounds[3]+5)))
    cropped.thumbnail((width - 24, height - 24), Image.Resampling.LANCZOS)
    normalized = Image.new('RGB', (width, height), 'white')
    normalized.paste(cropped, ((width-cropped.width)//2, (height-cropped.height)//2))
    output = io.BytesIO()
    normalized.save(output, format='PNG', optimize=True)
    encoded = base64.b64encode(output.getvalue()).decode('ascii')
    if len(encoded) > 48000:
        raise ValueError('簽名資料超過 Google Sheets 單格限制。')
    return encoded


def signature_from_upload(uploaded):
    image = Image.open(uploaded).convert('RGB')
    image.thumbnail((550, 190))
    white = Image.new('RGB', (550, 190), 'white')
    white.paste(image, ((550-image.width)//2, (190-image.height)//2))
    output = io.BytesIO()
    white.save(output, format='PNG', optimize=True)
    encoded = base64.b64encode(output.getvalue()).decode('ascii')
    if len(encoded) > 48000:
        output = io.BytesIO()
        white.convert('L').point(lambda x: 255 if x > 190 else 0, '1').save(output, format='PNG', optimize=True)
        encoded = base64.b64encode(output.getvalue()).decode('ascii')
    if len(encoded) > 48000:
        raise ValueError('簽名圖片過大，請使用白底黑字簽名圖片。')
    return encoded


def make_ticket_pdf(ticket):
    """保留紙本售後服務單 A4 單頁格式。"""
    font = 'MSung-Light'
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    output = io.BytesIO()
    c = pdf_canvas.Canvas(output, pagesize=A4)
    c.setTitle('JENN WEI MACHINERY - 售後服務單')
    pw, ph = A4
    x0, x1 = 15*mm, pw-15*mm
    W = x1-x0
    ink, red = colors.HexColor('#242424'), colors.HexColor('#9B352E')
    c.setStrokeColor(ink)
    c.setLineWidth(.65)

    def txt(x, y, value, size=9, color=ink):
        c.setFillColor(color)
        c.setFont(font, size)
        c.drawString(x, y, str(value))

    def value(key):
        return safe_str(ticket.get(key, ''))

    def cell(x, y, w, h, label, content='', label_width=None, fontsize=9):
        c.rect(x, y, w, h)
        txt(x+4, y+h-13, label, 9, red)
        if content:
            label_width = label_width if label_width is not None else min(w*.36, 31*mm)
            p = Paragraph(escape(str(content)).replace('\n', '<br/>'),
                          ParagraphStyle('v', fontName=font, fontSize=fontsize, leading=fontsize+2, wordWrap='CJK'))
            _, ht = p.wrap(max(12, w-label_width-7), 1000)
            if ht <= h-4:
                p.drawOn(c, x+label_width, y+h-3-ht)
            else:
                txt(x+label_width, y+h-13, str(content)[:18]+'…', 7)

    def multiline(x, y, w, h, label, content):
        c.rect(x, y, w, h)
        txt(x+5, y+h-14, label, 9, red)
        for size in (9, 8, 7, 6, 5):
            p = Paragraph(escape(str(content or '')).replace('\n', '<br/>') or ' ',
                          ParagraphStyle('m', fontName=font, fontSize=size, leading=size+2, wordWrap='CJK'))
            _, ht = p.wrap(w-13, 10000)
            if ht <= h-25:
                p.drawOn(c, x+6, y+h-22-ht)
                return
        c.saveState()
        scale = min(1, (h-25)/max(ht, 1))
        c.translate(x+6, y+4)
        c.scale(scale, scale)
        p.drawOn(c, 0, 0)
        c.restoreState()

    logo = os.path.join(os.path.dirname(__file__), 'jenn_wei_logo.png')
    if os.path.exists(logo):
        try:
            c.drawImage(logo, x0, ph-48*mm, width=36*mm, height=32*mm, preserveAspectRatio=True, mask='auto')
        except Exception:
            pass
    txt(x0+41*mm, ph-20*mm, '震唯機械股份有限公司', 15)
    c.setFont('Helvetica-Bold', 10)
    c.drawString(x0+41*mm, ph-25*mm, 'JENN WEI MACHINERY CO., LTD.')
    c.setFont('Helvetica', 7.5)
    for offset, label in [(30, 'E-mail: L3352368@ms49.hinet.net'), (34, 'http://www.jennwei.com.tw/'),
                          (38, 'TEL: 886-4-23352368   FAX: 886-4-23353880')]:
        c.drawString(x0+41*mm, ph-offset*mm, label)
    c.setFillColor(red)
    c.setFont(font, 14)
    c.drawCentredString(pw/2, ph-53*mm, '售後服務單')
    c.setFillColor(ink)
    c.setFont('Helvetica', 8)
    c.drawRightString(x1, ph-53*mm, 'No. '+value('ticket_id'))
    row = 11*mm
    y = ph-59*mm-row
    cell(x0, y, W*.46, row, '客戶名稱：', value('customer'))
    cell(x0+W*.46, y, W*.31, row, '電話：', value('phone'))
    cell(x0+W*.77, y, W*.23, row, '日期：', value('created_at')[:10], 18*mm, 8)
    y -= row
    cell(x0, y, W, row, '住址：', value('address'))
    y -= row
    cell(x0, y, W*.35, row, '機型：', value('machine_model'))
    cell(x0+W*.35, y, W*.35, row, '機號：', value('serial_number'))
    cell(x0+W*.70, y, W*.30, row, '出廠日：', value('manufacture_date'))
    y -= 29*mm
    multiline(x0, y, W, 29*mm, '故障原因：', value('diagnosis') or value('issue'))
    y -= 30*mm
    multiline(x0, y, W, 30*mm, '處理方式：', value('repair_action') or value('resolution'))
    y -= 10*mm
    cell(x0, y, W*.30, 10*mm, '到達時間：', value('arrival_time'))
    cell(x0+W*.30, y, W*.30, 10*mm, '離開時間：', value('departure_time'))
    completed = '是' if value('is_completed') == '是' or (not value('is_completed') and value('status') == '已結案') else '待確認'
    cell(x0+W*.60, y, W*.40, 10*mm, '本案是否完成：', completed, 33*mm)
    y -= 10*mm
    cell(x0, y, W*.76, 10*mm, '未處理完成、再處理日期：', value('next_service_date') or value('scheduled_date'))
    cell(x0+W*.76, y, W*.24, 10*mm, '時 間：', value('next_service_hours'), 17*mm)
    parts_h = 6.4*mm
    y -= parts_h
    fractions = [.50, .14, .14, .22]
    cursor = x0
    for frac, head in zip(fractions, ['零件名稱編號', '單價', '數量', '小計']):
        c.rect(cursor, y, W*frac, parts_h)
        c.setFont(font, 8)
        c.drawCentredString(cursor+W*frac/2, y+2*mm, head)
        cursor += W*frac
    legacy = [p.strip() for p in value('parts').replace('；', '\n').replace(';', '\n').splitlines() if p.strip()]
    total_parts = 0.0
    displayed = 0
    for n in range(1, 11):
        name = value(f'part_{n}_name') or (legacy[n-1] if n <= len(legacy) else '')
        if not name.strip():
            continue
        displayed += 1
        price, qty = number(ticket.get(f'part_{n}_price')), number(ticket.get(f'part_{n}_qty'))
        subtotal = price*qty
        total_parts += subtotal
        y -= parts_h
        cursor = x0
        entries = [f'{displayed}. {name}', f'{price:,.2f}' if price else '', f'{qty:g}' if qty else '', f'{subtotal:,.2f}' if price and qty else '']
        for frac, entry in zip(fractions, entries):
            c.rect(cursor, y, W*frac, parts_h)
            if entry:
                txt(cursor+3, y+2*mm, entry[:46], 7)
            cursor += W*frac
    y -= 10*mm
    fee, tax = number(ticket.get('service_fee')), number(ticket.get('tax_amount'))
    cell(x0, y, W*.34, 10*mm, '服務費：', f'{fee:,.2f}')
    cell(x0+W*.34, y, W*.30, 10*mm, '5% 稅金：', f'{tax:,.2f}')
    cell(x0+W*.64, y, W*.36, 10*mm, '費用總計：', f'{total_parts+fee+tax:,.2f}')
    y -= 11*mm
    cell(x0, y, W*.33, 11*mm, '服務人員：', value('engineer'))
    cell(x0+W*.33, y, W*.33, 11*mm, '售服主管：', value('service_supervisor'))
    cell(x0+W*.66, y, W*.34, 11*mm, '業務主管：', value('sales_supervisor'))
    sign_h = 24*mm
    y -= sign_h
    c.rect(x0, y, W, sign_h)
    txt(x0+4, y+sign_h-11, '客戶確認／簽名：', 9, red)
    if value('customer_signature'):
        try:
            source = Image.open(io.BytesIO(base64.b64decode(value('customer_signature')))).convert('RGB')
            bounds = source.convert('L').point(lambda px: 255 if px < 205 else 0).getbbox()
            if bounds:
                source = source.crop(bounds)
                factor = min(72*mm/source.width, 15*mm/source.height, 72/110)
                iw, ih = source.width*factor, source.height*factor
                c.drawImage(ImageReader(source), x0+42*mm+(W-47*mm-iw)/2,
                            y+(sign_h-ih)/2-1*mm, width=iw, height=ih)
        except Exception:
            txt(x0+45*mm, y+10*mm, '簽名圖片無法載入', 8)
    txt(x0+3, y-10, '簽收人：'+value('signed_by')+'   簽收時間：'+value('signed_at'), 8)
    txt(x0+3, y-23, '本單據記錄維修及簽收資訊；簽收不代表另行承諾保固或費用條件。', 7)
    c.showPage()
    c.save()
    return output.getvalue()


# Excel 匯出函式直接內建，不需要額外的 cc_excel_export.py。
"""Excel export for JENN-WEI cc.py. Uses existing openpyxl and Pillow dependencies."""
import base64
import io
import math
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.worksheet.page import PageMargins
from PIL import Image


def make_ticket_excel(ticket):
    def text(key):
        value = ticket.get(key, '')
        if value is None or str(value).lower() in ('nan', 'nat', '<na>'):
            return ''
        return str(value)

    def money(key):
        try:
            n = Decimal(str(ticket.get(key, '') or 0))
            return n if n.is_finite() else Decimal(0)
        except Exception:
            return Decimal(0)

    def rounded(n):
        return float(n.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))

    wb = Workbook()
    ws = wb.active
    ws.title = '售後服務單'
    widths = {'A': 14, 'B': 18, 'C': 18, 'D': 15, 'E': 15, 'F': 20}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    edge = Side(style='thin', color='AAB5C3')

    def block(row, start, end, value, title=False, height=28):
        for col in range(start, end + 1):
            c = ws.cell(row, col)
            c.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            c.alignment = Alignment(vertical='center', wrap_text=True)
            c.font = Font(name='Microsoft JhengHei', size=11, bold=title,
                          color='FFFFFF' if title else '242424')
            if title:
                c.fill = PatternFill('solid', fgColor='243D58')
        if end > start:
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        c = ws.cell(row, start)
        c.value = str(value)
        c.data_type = 's'  # Customer text must never become an Excel formula.
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, height)

    block(1, 1, 6, '震唯機械股份有限公司', True, 32)
    block(2, 1, 6, 'JENN WEI MACHINERY CO., LTD.｜售後服務單', False, 30)
    block(3, 1, 6, '工單號碼：' + text('ticket_id'))
    block(4, 1, 4, '客戶名稱：' + text('customer'))
    block(4, 5, 6, '日期：' + text('created_at')[:10])
    block(5, 1, 3, '聯絡人：' + text('contact'))
    block(5, 4, 6, '電話：' + text('phone'))
    block(6, 1, 6, '住址：' + text('address'))
    block(7, 1, 3, '機型：' + text('machine_model'))
    block(7, 4, 6, '機號：' + text('serial_number'))
    block(8, 1, 3, '出廠日：' + text('manufacture_date'))
    block(8, 4, 6, '保固：' + text('warranty'))
    for row, label, content in [(9, '故障原因', text('diagnosis') or text('issue')),
                                (10, '處理方式', text('repair_action') or text('resolution'))]:
        estimated_lines = sum(max(1, math.ceil(len(line) / 48)) for line in content.split('\n'))
        block(row, 1, 6, label + '：\n' + content, height=max(70, (estimated_lines + 1) * 17))
    block(11, 1, 3, '到達時間：' + text('arrival_time'))
    block(11, 4, 6, '離開時間：' + text('departure_time'))
    completed = text('is_completed') or ('是' if text('status') == '已結案' else '待確認')
    block(12, 1, 3, '本案是否完成：' + completed)
    block(12, 4, 6, '狀態：' + text('status'))
    block(13, 1, 4, '再處理日期：' + text('next_service_date'))
    block(13, 5, 6, '時間：' + text('next_service_hours'))
    block(14, 1, 3, '零件名稱／編號', True)
    for col, name in [(4, '單價'), (5, '數量'), (6, '小計')]:
        block(14, col, col, name, True)
    legacy = text('parts').replace('；', '\n').replace(';', '\n').splitlines()
    total_parts = Decimal(0)
    for n in range(1, 11):
        row = 14 + n
        name = text(f'part_{n}_name') or (legacy[n - 1].strip() if n <= len(legacy) else '')
        block(row, 1, 3, f'{n}. {name}' if name else '')
        price, qty = money(f'part_{n}_price'), money(f'part_{n}_qty')
        for col in (4, 5, 6):
            block(row, col, col, '')
        if name:
            subtotal = price * qty
            total_parts += subtotal
            for col, number in [(4, rounded(price)), (5, float(qty)), (6, rounded(subtotal))]:
                ws.cell(row, col, number)
                ws.cell(row, col).number_format = '#,##0.00'
    for row, label, value in [(25, '零件合計', total_parts),
                               (26, '服務費', money('service_fee')),
                               (27, '稅金（已填寫金額）', money('tax_amount')),
                               (28, '費用總計', total_parts + money('service_fee') + money('tax_amount'))]:
        block(row, 1, 5, label, row == 28)
        block(row, 6, 6, '', row == 28)
        ws.cell(row, 6, rounded(value)).number_format = '#,##0.00'
    block(29, 1, 2, '服務人員：' + text('engineer'))
    block(29, 3, 4, '售服主管：' + text('service_supervisor'))
    block(29, 5, 6, '業務主管：' + text('sales_supervisor'))
    block(30, 1, 6, '客戶確認／簽名：', height=100)
    signature = text('customer_signature')
    if signature:
        try:
            image = Image.open(io.BytesIO(base64.b64decode(signature))).convert('RGB')
            image.thumbnail((320, 105))
            stream = io.BytesIO()
            image.save(stream, format='PNG')
            stream.seek(0)
            picture = ExcelImage(stream)
            ws.add_image(picture, 'C30')
        except Exception:
            block(30, 1, 6, '客戶確認／簽名：圖片無法載入，請檢查已儲存簽名。', height=100)
    block(31, 1, 3, '簽收人：' + text('signed_by'))
    block(31, 4, 6, '簽收時間：' + text('signed_at'))
    block(32, 1, 6, '本單據記錄維修及簽收資訊；簽收不代表另行承諾保固或費用條件。')
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.print_area = 'A1:F32'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0  # Long descriptions may use multiple pages; preserve legibility.
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=.3, right=.3, top=.35, bottom=.35)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


df = read_tickets()
EMPLOYEE_NAMES = employee_names()
new_tab, manage_tab, form_tab, dashboard_tab, history_tab = st.tabs(
    ['📝 建立報修單', '🔧 維修工單管理', '📄 售後服務單填寫', '📊 售服 KPI', '📚 維修紀錄'])

with new_tab:
    st.subheader('建立客戶報修工單')
    from shared_contacts import customer_options
    customer_labels, customer_records = customer_options()
    selected_customer = st.selectbox('從共用主檔選擇客戶', ['手動輸入']+customer_labels, key='service_customer_choice')
    customer_record = customer_records.get(selected_customer, {})
    with st.form('new_service_ticket', clear_on_submit=True):
        a, b, c = st.columns(3)
        with a:
            customer = st.text_input('客戶名稱 *', value=str(customer_record.get('company_name', '')))
            contact = st.text_input('聯絡人', value=str(customer_record.get('contact', '')))
            phone = st.text_input('電話 / 聯絡方式', value=str(customer_record.get('phone', '')))
        with b:
            model = st.text_input('機台型號 *')
            serial = st.text_input('機台序號')
            warranty = st.selectbox('保固狀態', WARRANTIES)
        with c:
            priority = st.selectbox('緊急程度', PRIORITIES)
            engineer = employee_selector('預定負責工程師', 'new_service_engineer', names=EMPLOYEE_NAMES)
            schedule = st.date_input('預計處理日期', value=date.today())
        st.markdown('#### 售後服務單基本資料')
        address = st.text_input('客戶住址', value=str(customer_record.get('address', '')))
        manufacture_date = st.text_input('機台出廠日（YYYY-MM-DD，可留空）')
        issue = st.text_area('故障現象 / 客戶反映 *')
        submitted = st.form_submit_button('➕ 建立維修工單', type='primary')
    if submitted:
        if not customer.strip() or not model.strip() or not issue.strip():
            st.error('請填寫客戶名稱、機台型號及故障現象。')
        else:
            now = now_text()
            ticket = {col: '' for col in COLUMNS}
            ticket.update(ticket_id=f'SRV-{datetime.now():%Y%m%d}-{uuid.uuid4().hex[:4].upper()}',
                          created_at=now, customer=customer.strip(), contact=contact.strip(), phone=phone.strip(),
                          machine_model=model.strip(), serial_number=serial.strip(), warranty=warranty,
                          issue=issue.strip(), priority=priority, status='新報修', engineer=engineer.strip(),
                          scheduled_date=str(schedule), parts_cost=0.0, labor_hours=0.0, labor_rate=0.0,
                          total_cost=0.0, updated_at=now, address=address.strip(), manufacture_date=manufacture_date.strip())
            save_tickets(pd.concat([df, pd.DataFrame([ticket])], ignore_index=True))

with manage_tab:
    st.subheader('維修工單管理')
    a, b, c = st.columns(3)
    with a:
        status_filter = st.selectbox('維修狀態', ['全部']+STATUSES)
    with b:
        engineer_filter = st.selectbox('負責工程師', ['全部']+sorted(x for x in df['engineer'].dropna().astype(str).unique() if x.strip()))
    with c:
        keyword = st.text_input('搜尋工單 / 客戶 / 型號 / 序號')
    view = df.copy()
    if status_filter != '全部':
        view = view[view['status'] == status_filter]
    if engineer_filter != '全部':
        view = view[view['engineer'] == engineer_filter]
    if keyword:
        view = view[view.astype(str).apply(lambda s: s.str.contains(keyword, case=False, regex=False, na=False)).any(axis=1)]
    st.caption(f'符合條件工單：{len(view)} 筆')
    if view.empty:
        st.info('目前沒有符合條件的維修工單。')
    else:
        st.dataframe(view[['ticket_id', 'customer', 'machine_model', 'serial_number', 'priority', 'status', 'engineer', 'scheduled_date', 'updated_at']], hide_index=True, use_container_width=True)
        ticket_id = st.selectbox('選擇要更新的工單', view['ticket_id'].astype(str).tolist())
        idx = df.index[df['ticket_id'].astype(str) == ticket_id][0]
        item = df.loc[idx]
        with st.form(f'edit_service_ticket_{ticket_id}'):
            st.markdown(f'**{ticket_id}｜{safe_str(item["customer"])}｜{safe_str(item["machine_model"])}**')
            a, b, c = st.columns(3)
            with a:
                current = safe_str(item['status'])
                status = st.selectbox('維修狀態', STATUSES, index=STATUSES.index(current) if current in STATUSES else 0)
                current_priority = safe_str(item['priority'])
                priority_new = st.selectbox('優先程度', PRIORITIES, index=PRIORITIES.index(current_priority) if current_priority in PRIORITIES else 0)
            with b:
                engineer_new = employee_selector('負責工程師', f'edit_engineer_{ticket_id}', safe_str(item['engineer']), EMPLOYEE_NAMES)
                current_warranty = safe_str(item['warranty'])
                warranty_new = st.selectbox('保固', WARRANTIES, index=WARRANTIES.index(current_warranty) if current_warranty in WARRANTIES else 0)
            with c:
                parts_cost = st.number_input('零件費用', min_value=0.0, value=number(item['parts_cost']))
                labor_hours = st.number_input('維修工時（小時）', min_value=0.0, value=number(item['labor_hours']))
                labor_rate = st.number_input('每小時工資', min_value=0.0, value=number(item['labor_rate']))
            diagnosis = st.text_area('故障診斷', value=safe_str(item['diagnosis']))
            action = st.text_area('維修處置', value=safe_str(item['repair_action']))
            parts = st.text_area('更換零件', value=safe_str(item['parts']))
            resolution = st.text_area('測試 / 客戶驗收紀錄', value=safe_str(item['resolution']))
            st.metric('本次維修費用合計', f'{parts_cost+labor_hours*labor_rate:,.0f}')
            saved = st.form_submit_button('💾 儲存維修紀錄', type='primary')
        if saved:
            updates = dict(status=status, priority=priority_new, engineer=engineer_new.strip(), warranty=warranty_new,
                           parts_cost=parts_cost, labor_hours=labor_hours, labor_rate=labor_rate,
                           total_cost=parts_cost+labor_hours*labor_rate, diagnosis=diagnosis,
                           repair_action=action, parts=parts, resolution=resolution, updated_at=now_text())
            updates['closed_at'] = (safe_str(item['closed_at']) or now_text()) if status == '已結案' else ''
            for key, val in updates.items():
                df.at[idx, key] = val
            save_tickets(df)

with form_tab:
    st.subheader('📄 售後服務單 — 對應紙本欄位填寫')
    st.caption('依紙本欄位填寫，儲存後可下載 Excel 或 PDF（含已儲存簽名）。')
    if df.empty:
        st.info('請先在「建立報修單」建立工單。')
    else:
        form_ticket_id = st.selectbox('選擇維修單', df['ticket_id'].astype(str).tolist(), key='form_ticket_id')
        form_idx = df.index[df['ticket_id'].astype(str) == form_ticket_id][0]
        item = df.loc[form_idx]

        def old(key):
            return safe_str(item.get(key, ''))

        count_key = f'parts_visible_count_{form_ticket_id}'
        existing_slots = [n for n in range(1, 11) if old(f'part_{n}_name').strip()]
        if count_key not in st.session_state:
            st.session_state[count_key] = max(1, max(existing_slots, default=0))
        st.caption(f'零件明細：目前顯示 {st.session_state[count_key]} 筆，最多 10 筆。')
        if st.button('➕ 追加一筆零件', key=f'add_part_{form_ticket_id}', disabled=st.session_state[count_key] >= 10):
            st.session_state[count_key] += 1
            st.rerun()
        st.markdown('#### 1. 客戶及機台資料')
        a, b, c = st.columns(3)
        with a:
            customer_f = st.text_input('客戶名稱', value=old('customer'), key=f'customer_{form_ticket_id}')
            address_f = st.text_input('客戶住址', value=old('address'), key=f'address_{form_ticket_id}')
            machine_f = st.text_input('機型', value=old('machine_model'), key=f'model_{form_ticket_id}')
        with b:
            phone_f = st.text_input('電話', value=old('phone'), key=f'phone_{form_ticket_id}')
            serial_f = st.text_input('機號', value=old('serial_number'), key=f'serial_{form_ticket_id}')
            manufactured_f = st.text_input('出廠日', value=old('manufacture_date'), placeholder='YYYY-MM-DD', key=f'manufacture_{form_ticket_id}')
        with c:
            created_f = st.text_input('服務單日期', value=old('created_at')[:10], placeholder='YYYY-MM-DD', key=f'created_{form_ticket_id}')
            engineer_f = employee_selector('服務人員', f'service_form_engineer_{form_ticket_id}', old('engineer'), EMPLOYEE_NAMES)
        st.markdown('#### 2. 故障原因與處理方式（按需追加）')

        def split_entries(raw):
            import re
            return [clean for line in safe_str(raw).splitlines()
                    if (clean := re.sub(r'^\s*(?:\d+[.、．)]\s*)', '', line).strip())]

        issue_values = split_entries(old('diagnosis') or old('issue'))
        action_values = split_entries(old('repair_action'))
        issue_count_key, action_count_key = f'issue_count_{form_ticket_id}', f'action_count_{form_ticket_id}'
        if issue_count_key not in st.session_state:
            st.session_state[issue_count_key] = max(1, len(issue_values))
        if action_count_key not in st.session_state:
            st.session_state[action_count_key] = max(1, len(action_values))
        st.markdown('**故障原因**')
        if st.button('➕ 追加故障原因', key=f'add_issue_{form_ticket_id}'):
            st.session_state[issue_count_key] += 1
            st.rerun()
        diagnosis_entries = [st.text_area(f'故障原因 {n+1}', value=issue_values[n] if n < len(issue_values) else '',
                                          height=80, key=f'issue_{form_ticket_id}_{n}')
                             for n in range(st.session_state[issue_count_key])]
        st.markdown('**處理方式**')
        if st.button('➕ 追加處理方式', key=f'add_action_{form_ticket_id}'):
            st.session_state[action_count_key] += 1
            st.rerun()
        repair_entries = [st.text_area(f'處理方式 {n+1}', value=action_values[n] if n < len(action_values) else '',
                                       height=80, key=f'action_{form_ticket_id}_{n}')
                          for n in range(st.session_state[action_count_key])]
        st.markdown('#### 3. 服務時間與完工狀態')
        a, b, c = st.columns(3)
        with a:
            arrival_f = st.text_input('到達時間', value=old('arrival_time'), placeholder='YYYY-MM-DD HH:MM', key=f'arrival_{form_ticket_id}')
        with b:
            departure_f = st.text_input('離開時間', value=old('departure_time'), placeholder='YYYY-MM-DD HH:MM', key=f'departure_{form_ticket_id}')
        with c:
            completed_default = old('is_completed') or ('是' if old('status') == '已結案' else '否')
            completed_f = st.radio('本案是否完成', ['是', '否'], index=0 if completed_default == '是' else 1,
                                   horizontal=True, key=f'completed_{form_ticket_id}')
        a, b = st.columns(2)
        with a:
            next_date_f = st.text_input('未完成，再處理日期', value=old('next_service_date'), placeholder='YYYY-MM-DD', key=f'nextdate_{form_ticket_id}')
        with b:
            next_hours_f = st.text_input('再處理時間', value=old('next_service_hours'), placeholder='HH:MM', key=f'nexthours_{form_ticket_id}')
        st.markdown('#### 4. 零件名稱、單價、數量（按需追加）')
        parts_input = []
        for n in range(1, st.session_state[count_key]+1):
            a, b, c = st.columns([5, 2, 2])
            with a:
                name_f = st.text_input(f'{n}. 零件名稱／編號', value=old(f'part_{n}_name'), key=f'partname_{form_ticket_id}_{n}')
            with b:
                price_f = st.number_input(f'{n}. 單價', min_value=0.0, value=number(item.get(f'part_{n}_price', '')), step=1.0, key=f'partprice_{form_ticket_id}_{n}')
            with c:
                qty_f = st.number_input(f'{n}. 數量', min_value=0.0, value=number(item.get(f'part_{n}_qty', '')), step=1.0, key=f'partqty_{form_ticket_id}_{n}')
            parts_input.append((name_f, price_f, qty_f))
        st.markdown('#### 5. 費用與主管確認')
        a, b, c = st.columns(3)
        with a:
            service_fee_f = st.number_input('服務費', min_value=0.0, value=number(item.get('service_fee', '')), step=100.0, key=f'fee_{form_ticket_id}')
        with b:
            tax_f = st.number_input('5% 稅金（手動填寫）', min_value=0.0, value=number(item.get('tax_amount', '')), step=1.0, key=f'tax_{form_ticket_id}')
        with c:
            subtotal_f = sum(price*qty for name, price, qty in parts_input if name.strip())
            st.metric('費用總計（零件＋服務費＋稅金）', f'{subtotal_f+service_fee_f+tax_f:,.2f}')
        a, b = st.columns(2)
        with a:
            supervisor_f = employee_selector('售服主管', f'service_form_supervisor_{form_ticket_id}', old('service_supervisor'), EMPLOYEE_NAMES)
        with b:
            sales_f = employee_selector('業務主管', f'service_form_sales_{form_ticket_id}', old('sales_supervisor'), EMPLOYEE_NAMES)
        if st.button('💾 儲存售後服務單資料', type='primary', use_container_width=True, key=f'save_paper_{form_ticket_id}'):
            updates = {'customer': customer_f.strip(), 'address': address_f.strip(), 'phone': phone_f.strip(),
                       'machine_model': machine_f.strip(), 'serial_number': serial_f.strip(),
                       'manufacture_date': manufactured_f.strip(), 'engineer': engineer_f.strip(),
                       'diagnosis': '\n'.join(f'{n}. {v.strip()}' for n, v in enumerate((x for x in diagnosis_entries if x.strip()), 1)),
                       'repair_action': '\n'.join(f'{n}. {v.strip()}' for n, v in enumerate((x for x in repair_entries if x.strip()), 1)),
                       'arrival_time': arrival_f.strip(), 'departure_time': departure_f.strip(), 'is_completed': completed_f,
                       'next_service_date': next_date_f.strip(), 'next_service_hours': next_hours_f.strip(),
                       'service_fee': service_fee_f, 'tax_amount': tax_f, 'service_supervisor': supervisor_f.strip(),
                       'sales_supervisor': sales_f.strip(), 'parts_cost': subtotal_f,
                       'total_cost': subtotal_f+service_fee_f+tax_f, 'updated_at': now_text()}
            if created_f.strip():
                original_time = old('created_at')
                updates['created_at'] = created_f.strip()+(' '+original_time.split(' ', 1)[1] if ' ' in original_time else '')
            for n, (name, price, qty) in enumerate(parts_input, 1):
                updates.update({f'part_{n}_name': name.strip(), f'part_{n}_price': price, f'part_{n}_qty': qty})
            for key, val in updates.items():
                df.at[form_idx, key] = val
            save_tickets(df)

    st.divider()
    st.subheader('✍️ 客戶手機簽名與維修單 Excel／PDF')
    st.caption('簽名儲存至同一份 Google Sheets 的 ServiceTickets 分頁。')
    if df.empty:
        st.info('請先建立維修工單。')
    else:
        ids = df['ticket_id'].astype(str).tolist()
        selected_id = st.selectbox('選擇維修單號', ids, index=ids.index(form_ticket_id), key='sign_ticket')
        selected_idx = df.index[df['ticket_id'].astype(str) == selected_id][0]
        ticket = df.loc[selected_idx]
        st.write(f'**客戶：** {safe_str(ticket["customer"])}　**機台：** {safe_str(ticket["machine_model"])}')
        if safe_str(ticket.get('customer_signature', '')):
            st.success(f'已簽收：{safe_str(ticket.get("signed_by", ""))}　{safe_str(ticket.get("signed_at", ""))}')
            try:
                st.image(base64.b64decode(safe_str(ticket['customer_signature'])), width=340)
            except Exception:
                st.warning('已存簽名資料，但無法顯示預覽。')
        else:
            st.info('本工單尚未簽名。')
        signer = st.text_input('客戶簽收人姓名', value=safe_str(ticket.get('signed_by', '')), key=f'signer_{selected_id}')
        drawing_canvas = None
        if st_canvas is None:
            st.warning('尚未安裝畫布元件，可使用下方簽名圖片上傳；手寫畫布需 streamlit-drawable-canvas。')
        else:
            st.caption('請在白色區域完整簽名。簽名區：340 × 170 px。')
            drawing_canvas = st_canvas(fill_color='rgba(255,255,255,0)', stroke_width=3, stroke_color='#17263C',
                                       background_color='#FFFFFF', update_streamlit=True, height=170, width=340,
                                       drawing_mode='freedraw', key=f'canvas_mobile_v2_{selected_id}')
        uploaded_signature = st.file_uploader('備用：上傳簽名圖片（PNG／JPG）', type=['png', 'jpg', 'jpeg'], key=f'upload_signature_{selected_id}')
        if st.button('💾 儲存客戶簽名', type='primary', key=f'save_signature_{selected_id}'):
            if not signer.strip():
                st.error('請先輸入簽收人姓名。')
            else:
                try:
                    if uploaded_signature is not None:
                        encoded = signature_from_upload(uploaded_signature)
                    elif drawing_canvas is not None:
                        encoded = signature_from_drawing(drawing_canvas.json_data or {}, width=340, height=170)
                    else:
                        raise ValueError('請上傳簽名圖片。')
                except Exception as exc:
                    st.error(f'簽名儲存失敗：{exc}')
                else:
                    df.at[selected_idx, 'customer_signature'] = encoded
                    df.at[selected_idx, 'signed_by'] = signer.strip()
                    df.at[selected_idx, 'signed_at'] = now_text()
                    df.at[selected_idx, 'updated_at'] = now_text()
                    save_tickets(df)
        st.divider()
        excel_col, pdf_col = st.columns(2)
        with excel_col:
            st.download_button('📊 下載維修服務單 Excel（含已儲存簽名）', data=make_ticket_excel(ticket),
                               file_name=f'{selected_id}_service_report.xlsx',
                               mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                               use_container_width=True, key=f'excel_download_{selected_id}')
        with pdf_col:
            st.download_button('📄 下載維修服務單 PDF（含已儲存簽名）', data=make_ticket_pdf(ticket),
                               file_name=f'{selected_id}_service_report.pdf', mime='application/pdf',
                               use_container_width=True, key=f'pdf_download_{selected_id}')
        st.caption('先儲存服務單資料及簽名，頁面重新載入後再下載，才會包含最新內容。')

with dashboard_tab:
    st.subheader('售服管理 KPI')
    total = len(df)
    closed = int((df['status'] == '已結案').sum())
    urgent = int(df['priority'].isin(['急件', '緊急']).sum())
    a, b, c, d = st.columns(4)
    a.metric('累計報修', total)
    b.metric('處理中', total-closed)
    c.metric('已結案', closed)
    d.metric('急件 / 緊急', urgent)
    st.markdown('#### 維修狀態分布')
    st.bar_chart(df['status'].value_counts().reindex(STATUSES, fill_value=0))
    st.markdown('#### 工程師工單數')
    if not df.empty:
        st.bar_chart(df['engineer'].replace('', '未指派').fillna('未指派').value_counts())
    st.metric('累計維修費用', f'{pd.to_numeric(df["total_cost"], errors="coerce").fillna(0).sum():,.0f}')

with history_tab:
    st.subheader('維修紀錄查詢 / 匯出')
    st.dataframe(df.sort_values('created_at', ascending=False), hide_index=True, use_container_width=True)
    st.download_button('⬇️ 匯出維修工單 CSV', data=df.to_csv(index=False).encode('utf-8-sig'),
                       file_name='service_tickets.csv', mime='text/csv')

with st.expander('⚙️ Google Sheets 工作表設定'):
    st.write(f'使用與 AA、BB 相同的 Google Spreadsheet，工作表名稱為 **{SHEET}**。第一列欄位如下：')
    st.code(','.join(COLUMNS))
    st.info('既有 ServiceTickets 分頁不需重建；請保留原資料，補齊缺少的簽名與售後服務單欄位。Excel／PDF 下載不需新增分頁。')
    st.caption('讀取快取 TTL 為 600 秒；寫入後清除工單快取。')
